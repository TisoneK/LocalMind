"""B-2026-08-23-6: the file-content preview endpoint.

GET /api/files/content returns readable text for any file the agent created —
text/code verbatim, office docs (docx/pdf/xlsx) parsed through the same
tools.file_reader.parse_file the engine uses, and images as inline data URLs —
so the UI can render created files without a binary round trip.
"""
from __future__ import annotations

import asyncio
import base64
from unittest.mock import patch

import pytest

import core.config as config_module
from api.routes.filesystem import get_file_content


def _run(coro):
    return asyncio.run(coro)


def _allow(tmp_path):
    """Scoped helper: allow the tmp dir so the endpoint passes its gate.

    Patches the module-level folder list (the pydantic settings model forbids
    attribute patching, so we go one level down to _allowed_user_folders).
    """
    patcher = patch.object(config_module, "_allowed_user_folders", return_value=[str(tmp_path)])
    patcher.start()
    return patcher


class TestFileContentPreview:
    def test_plain_text_file(self, tmp_path):
        f = tmp_path / "notes.txt"
        f.write_text("hello world", encoding="utf-8")
        patcher = _allow(tmp_path)
        try:
            result = _run(get_file_content(str(f)))
        finally:
            patcher.stop()
        assert result["name"] == "notes.txt"
        assert result["text"] == "hello world"
        assert result["truncated"] is False
        assert result["editable"] is True

    def test_markdown_and_code_verbatim(self, tmp_path):
        f = tmp_path / "app.py"
        f.write_text("def main():\n    return 1\n", encoding="utf-8")
        patcher = _allow(tmp_path)
        try:
            result = _run(get_file_content(str(f)))
        finally:
            patcher.stop()
        assert result["extension"] == "py"
        assert "def main()" in result["text"]
        assert result["editable"] is True

    def test_image_returns_data_url(self, tmp_path):
        f = tmp_path / "pic.png"
        f.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
        patcher = _allow(tmp_path)
        try:
            result = _run(get_file_content(str(f)))
        finally:
            patcher.stop()
        assert result["image_data_url"].startswith("data:image/png;base64,")
        # Round-trip: the decoded payload matches the file bytes.
        encoded = result["image_data_url"].split(",", 1)[1]
        assert base64.b64decode(encoded) == f.read_bytes()
        # Images are view-only — never editable.
        assert result["editable"] is False

    def test_office_doc_parsed_via_file_reader(self, tmp_path):
        f = tmp_path / "report.docx"
        f.write_bytes(b"fake docx bytes")
        patcher = _allow(tmp_path)

        class _FakeAttachment:
            chunks = ["Heading", "Body paragraph", "Closing"]

        try:
            with patch("tools.file_reader.parse_file", return_value=_FakeAttachment()) as mocked:
                result = _run(get_file_content(str(f)))
        finally:
            patcher.stop()
        assert mocked.called
        assert result["extension"] == "docx"
        assert "Heading\n\nBody paragraph" in result["text"]
        # Parsed office text is lossy — writing it back would destroy the
        # document, so it must never be marked editable.
        assert result["editable"] is False

    def test_xlsx_returns_structured_sheets(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["item", "qty", "price"])
        ws.append(["apple", 3, 1.5])
        ws.append(["pear", 2, 2.0])
        ws2 = wb.create_sheet("Notes")
        ws2.append(["note"])
        ws2.append(["hello"])
        f = tmp_path / "book.xlsx"
        wb.save(f)

        patcher = _allow(tmp_path)
        try:
            result = _run(get_file_content(str(f)))
        finally:
            patcher.stop()

        assert result["editable"] is False
        sheets = result["sheets"]
        assert [s["name"] for s in sheets] == ["Sales", "Notes"]
        assert sheets[0]["rows"][0] == ["item", "qty", "price"]
        assert sheets[0]["rows"][1] == ["apple", "3", "1.5"]
        assert sheets[1]["rows"][1] == ["hello"]

    def test_csv_returns_single_sheet(self, tmp_path):
        f = tmp_path / "data.csv"
        f.write_text("name,score\nalice,10\nbob,20\n", encoding="utf-8")
        patcher = _allow(tmp_path)
        try:
            result = _run(get_file_content(str(f)))
        finally:
            patcher.stop()

        assert result["editable"] is False
        assert [s["name"] for s in result["sheets"]] == ["Sheet 1"]
        rows = result["sheets"][0]["rows"]
        assert rows[0] == ["name", "score"]
        assert rows[1] == ["alice", "10"]
        assert rows[2] == ["bob", "20"]

    def test_xlsx_rows_capped_and_flagged(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["n"])
        for i in range(250):
            ws.append([i])
        f = tmp_path / "big.xlsx"
        wb.save(f)

        patcher = _allow(tmp_path)
        try:
            result = _run(get_file_content(str(f)))
        finally:
            patcher.stop()

        sheet = result["sheets"][0]
        assert sheet["truncated"] is True
        # Header + cap rows.
        assert len(sheet["rows"]) == 201

    def test_path_outside_allowed_folders_403(self, tmp_path):
        f = tmp_path / "secret.txt"
        f.write_text("nope", encoding="utf-8")
        patcher = patch.object(config_module, "_allowed_user_folders", return_value=[])
        patcher.start()
        try:
            with pytest.raises(Exception) as exc:
                _run(get_file_content(str(f)))
            assert exc.value.status_code == 403
        finally:
            patcher.stop()

    def test_missing_file_404(self, tmp_path):
        patcher = _allow(tmp_path)
        try:
            with pytest.raises(Exception) as exc:
                _run(get_file_content(str(tmp_path / "missing.txt")))
            assert exc.value.status_code == 404
        finally:
            patcher.stop()

    def test_directory_400(self, tmp_path):
        patcher = _allow(tmp_path)
        try:
            with pytest.raises(Exception) as exc:
                _run(get_file_content(str(tmp_path)))
            assert exc.value.status_code == 400
        finally:
            patcher.stop()
